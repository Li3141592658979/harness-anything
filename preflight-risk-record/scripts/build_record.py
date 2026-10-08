# -*- coding: utf-8 -*-
"""运航部值班干部航前风险提示记录表生成脚本。

流程：读取出港签到表(xlsx)的航班信息 → 从风险文件夹历史记录表(docx)按航线
匹配风险类型与提示内容（取该航线最新一次出现的记录）→ 复制最新模板重建数据行
→ 生成当日记录表(docx)。

用法示例：
    python build_record.py --signin C:\\path\\出港签到表.xlsx \
        --risk-dir "C:\\Users\\Administrator\\Desktop\\飞行资料\\风险"

参数：
    --signin   出港签到表 xlsx（必填）
    --risk-dir 风险文件夹（默认：桌面飞行资料\\风险）
    --out      输出文件路径（默认：risk-dir/运航部值班干部航前风险提示记录表{M.D}.docx）
    --date     日期 YYYY-MM-DD（默认从签到表标题解析）
    --sheet    指定工作表名（默认选标题含“广州”的工作表；未找到则用首个有表头的表）
    --backup   备份机组名单，逗号分隔（默认从签到表“备份机组”行读取）
"""
import argparse
import copy
import datetime as _dt
import os
import re
import shutil
import sys

import docx
from docx.table import _Row as Row
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DEFAULT_RISK_DIR = r"C:\Users\Administrator\Desktop\飞行资料\风险"
RISK_ITEMS = ["机组", "飞机", "航线", "天气", "机场", "标准", "程序", "其他"]
LINE1_ITEMS = {"机组", "飞机", "航线", "天气"}


def parse_args():
    ap = argparse.ArgumentParser()
    ap.add_argument("--signin", required=True, help="出港签到表 xlsx 路径")
    ap.add_argument("--risk-dir", default=DEFAULT_RISK_DIR)
    ap.add_argument("--out", default=None)
    ap.add_argument("--date", default=None, help="YYYY-MM-DD，默认从签到表标题解析")
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--backup", default=None)
    return ap.parse_args()


def parse_signin(path, wanted_sheet=None):
    """解析签到表：返回 (date, flights, backup_names, used_sheet)。"""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = None
    if wanted_sheet:
        sheet = wb[wanted_sheet]
    else:
        cands = [ws for ws in wb.worksheets if "广州" in ws.title]
        if cands:
            sheet = cands[0]
        else:
            for ws in wb.worksheets:
                rows = list(ws.iter_rows(values_only=True))
                if any(rows) and any("航班" in str(v) for r in rows for v in r if v):
                    sheet = ws
                    break
    if sheet is None:
        raise RuntimeError("找不到可用的签到表工作表")

    rows = list(sheet.iter_rows(values_only=True))
    date = None
    for r in rows[:3]:
        for v in r:
            m = re.search(r"(\d{4})年(\d{1,2})月(\d{1,2})日", str(v or ""))
            if m:
                date = _dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
                break
        if date:
            break
    if date is None:
        raise RuntimeError("签到表标题中未找到日期（YYYY年M月D日）")

    # 定位表头行
    head_idx = None
    for i, r in enumerate(rows[:5]):
        vals = [str(v or "") for v in r]
        if any("航班" in v for v in vals) and any("航程" in v or "航线" in v for v in vals):
            head_idx = i
            break
    if head_idx is None:
        raise RuntimeError("未找到表头行（航班/航程）")

    def col(keywords):
        vals = [str(v or "") for v in rows[head_idx]]
        for k in keywords:
            for j, v in enumerate(vals):
                if k in v:
                    return j
        return None

    c_flight, c_route = col(["航班"]), col(["航程", "航线"])
    c_crew = col(["当日机组", "机组"])
    if c_flight is None or c_route is None:
        raise RuntimeError("表头缺少航班/航程列")

    flights, backup = [], []
    in_backup = False
    for r in rows[head_idx + 1:]:
        if not any(r):
            continue
        a = str(r[0] or "")
        if in_backup:
            if c_crew is not None and r[c_crew]:
                for n in re.split(r"[/、,，\s]+", str(r[c_crew]).strip()):
                    if n:
                        backup.append(n)
            continue
        if "备份" in a:
            in_backup = True
            if c_crew is not None and r[c_crew]:
                for n in re.split(r"[/、,，\s]+", str(r[c_crew]).strip()):
                    if n:
                        backup.append(n)
            continue
        if "值班" in a or "日期" in a:
            break
        flt = str(r[c_flight] or "").strip()
        route_raw = str(r[c_route] or "").strip()
        if not flt or not route_raw:
            continue
        numbers = [x.strip() for x in re.split(r"[/、,，]+", flt) if x.strip()]
        if numbers:
            m = re.match(r"^([A-Za-z]+)(\d.*)$", numbers[0])
            prefix = m.group(1) if m else ""
            numbers = [n if n[:2].isalpha() or not prefix else prefix + n for n in numbers]
        route = re.sub(r"\([A-Za-z0-9]+\)", "", route_raw).strip("- ")
        route = re.sub(r"\s+", "", route)
        flights.append((numbers, route))
    return date, flights, backup, sheet.title


def collect_risk_records(risk_dir):
    """读取风险文件夹全部 docx 表格，返回 (template_path, records)。

    records: [{route:[...], checked:set, contents:[...], mtime:float}]
    template 取表头匹配（航班号/航线/风险类型/风险提示内容）且 mtime 最新的文件。
    """
    template_path, template_mtime = None, -1
    records = []
    files = []
    for root, _, fnames in os.walk(risk_dir):
        for fn in fnames:
            if fn.lower().endswith(".docx"):
                files.append(os.path.join(root, fn))
    for fp in sorted(files):
        try:
            d = docx.Document(fp)
        except Exception:
            continue
        for tb in d.tables:
            if len(tb.rows) < 2:
                continue
            hdr = [c.text.strip() for c in tb.rows[0].cells]
            if not (any("航班" in h for h in hdr) and any("航线" in h or "航程" in h for h in hdr)
                    and any("风险" in h for h in hdr)):
                continue
            mtime = os.path.getmtime(fp)
            if mtime > template_mtime:
                template_path, template_mtime = fp, mtime
            for row in tb.rows[1:]:
                cells = [c.text for c in row.cells]
                if "备份" in cells[0] or (len(cells) > 1 and "备份" in cells[1]):
                    continue
                route = re.sub(r"\s+", "", cells[2].strip("- ")) if len(cells) > 2 else ""
                if not route or "航班" in route:
                    continue
                checked = set()
                risk_cell = cells[3] if len(cells) > 3 else ""
                for item in RISK_ITEMS:
                    if ("☑" + item) in risk_cell:
                        checked.add(item)
                contents = [p.strip() for p in
                            (cells[4].split("\n") if len(cells) > 4 else []) if p.strip()]
                records.append({
                    "route": [s for s in re.split(r"[-－—]", route) if s],
                    "checked": checked,
                    "contents": contents,
                    "mtime": mtime,
                    "source": os.path.basename(fp),
                })
    return template_path, records


def match_record(route_abbrevs, records):
    """精确航线优先，其次同起降点；均取 mtime 最新。无匹配返回 None。"""
    exact = [r for r in records if r["route"] == route_abbrevs]
    if exact:
        return max(exact, key=lambda r: r["mtime"])
    first, last = route_abbrevs[0], route_abbrevs[-1]
    cands = [r for r in records if r["route"] and r["route"][0] == first and r["route"][-1] == last]
    return max(cands, key=lambda r: r["mtime"]) if cands else None


def set_cell_texts(cell, texts):
    paras = cell.paragraphs
    while len(paras) < len(texts):
        new_p = copy.deepcopy(paras[-1]._p)
        paras[-1]._p.addnext(new_p)
        paras = cell.paragraphs
    while len(paras) > len(texts):
        p = paras.pop()
        p._p.getparent().remove(p._p)
        paras = cell.paragraphs
    for p, t in zip(cell.paragraphs, texts):
        runs = p.runs
        if runs:
            runs[0].text = t
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
        else:
            p.add_run(t)


def set_row_height(row, twips):
    trPr = row._tr.get_or_add_trPr()
    for el in trPr.findall(qn("w:trHeight")):
        trPr.remove(el)
    th = OxmlElement("w:trHeight")
    th.set(qn("w:val"), str(twips))
    th.set(qn("w:hRule"), "exact")
    trPr.append(th)


def est_height(flights, contents):
    visual = sum(max(1, (len(p) + 27) // 28) for p in contents)
    lines = max(len(flights), 2, visual)
    return 450 + 250 * (lines - 1)


def build(rows_data, template_path, out_path, date, backup_names):
    shutil.copy2(template_path, out_path)
    doc = docx.Document(out_path)
    tb = doc.tables[0]
    if len(tb.rows) < 3:
        raise RuntimeError("模板表格行数不足（需 表头+数据行+备份行）")

    proto_tr = copy.deepcopy(tb.rows[1]._tr)
    proto_row = Row(proto_tr, tb)
    risk1 = proto_row.cells[3].paragraphs[0].text.replace("☑", "□")
    risk2 = proto_row.cells[3].paragraphs[1].text.replace("☑", "□")

    for r in list(tb.rows[1:-1]):
        r._tr.getparent().remove(r._tr)

    header_tr = tb.rows[0]._tr
    last_tr = header_tr
    for flights, route, checked, contents in rows_data:
        tr = copy.deepcopy(proto_tr)
        last_tr.addnext(tr)
        last_tr = tr
        row = Row(tr, tb)
        set_cell_texts(row.cells[0], [""])
        set_cell_texts(row.cells[1], flights)
        set_cell_texts(row.cells[2], [route])
        l1, l2 = risk1, risk2
        for item in checked:
            if item in LINE1_ITEMS:
                l1 = l1.replace("□" + item, "☑" + item)
            else:
                l2 = l2.replace("□" + item, "☑" + item)
        set_cell_texts(row.cells[3], [l1, l2])
        set_cell_texts(row.cells[4], contents)
        set_cell_texts(row.cells[5], [""])
        set_row_height(row, est_height(flights, contents))

    if backup_names:
        backup_row = tb.rows[-1]
        set_cell_texts(backup_row.cells[3], backup_names)

    date_text = f"{date.year}年{date.month:02d}月{date.day:02d}日"
    replaced = False
    for p in doc.paragraphs:
        if "日期" in p.text:
            for r in p.runs:
                if re.search(r"\d{4}年\d{1,2}月\d{1,2}日", r.text):
                    r.text = re.sub(r"\d{4}年\d{1,2}月\d{1,2}日", date_text, r.text)
                    replaced = True
            if not replaced and p.runs:
                p.runs[0].text = re.sub(r"\d{4}年\d{1,2}月\d{1,2}日", date_text, p.text)
                for r in p.runs[1:]:
                    r._r.getparent().remove(r._r)
                replaced = True
    doc.save(out_path)
    return replaced


def main():
    args = parse_args()
    date, flights, backup, used_sheet = parse_signin(args.signin, args.sheet)
    if args.date:
        date = _dt.date.fromisoformat(args.date)
    if args.backup:
        backup = [n.strip() for n in re.split(r"[,，]", args.backup) if n.strip()]

    template_path, records = collect_risk_records(args.risk_dir)
    if template_path is None:
        sys.exit("风险文件夹中未找到含 航班号/航线/风险类型 表头的记录表模板")

    rows_data, unmatched = [], []
    for numbers, route_abbr in flights:
        rec = match_record(route_abbr.split("-"), records)
        if rec is None:
            unmatched.append((numbers, route_abbr))
            rows_data.append((numbers, route_abbr, set(), []))
        else:
            rows_data.append((numbers, route_abbr, rec["checked"], rec["contents"]))

    if args.out:
        out_path = args.out
    else:
        fname = f"运航部值班干部航前风险提示记录表{date.month}.{date.day}.docx"
        out_path = os.path.join(args.risk_dir, fname)

    replaced = build(rows_data, template_path, out_path, date, backup)

    print("date:", date, "| sheet:", used_sheet, "| flights:", len(flights),
          "| backup:", "、".join(backup) or "(空)")
    print("template:", os.path.basename(template_path))
    print("out:", out_path)
    print("date_replaced:", replaced)
    if unmatched:
        print("WARN 未匹配到历史风险的航班（风险内容留空）:")
        for n, r in unmatched:
            print("   ", "/".join(n), r)
    # 回读核对
    d2 = docx.Document(out_path)
    t2 = d2.tables[0]
    print("total_rows:", len(t2.rows))
    for i in (0, 1, len(t2.rows) // 2, len(t2.rows) - 2):
        print("ROW", i, "|", " || ".join(c.text.replace("\n", " / ") for c in t2.rows[i].cells)[:160])
    print("backup:", " || ".join(c.text.replace("\n", " / ") for c in t2.rows[-1].cells)[:80])


if __name__ == "__main__":
    main()
